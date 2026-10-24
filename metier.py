"""Règles métier : opérations, synthèses, import Excel contrôlé, sauvegarde Excel, prévisionnel."""
from __future__ import annotations

import hashlib
import io
import uuid
from datetime import date, datetime

import pandas as pd

from referentiel import (COLONNES_IMPORT, MANIFESTATION_GENERALE, MODES_PAIEMENT, SENS, TABLES,
                         categorie_a_classer, index_categories)


# ------------------------------------------------------------------ typage
def typer_operations(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    if df.empty:
        df = pd.DataFrame(columns=TABLES["operations"])
    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    df["montant"] = pd.to_numeric(df["montant"].astype(str).str.replace(",", "."), errors="coerce").fillna(0.0)
    df["exercice"] = pd.to_numeric(df["exercice"], errors="coerce").fillna(0).astype(int)
    df["signe"] = df["montant"].where(df["sens"] == "Recette", -df["montant"])
    return df


def exercice_de(d: date, mois_debut: int = 1) -> int:
    """Exercice = année de début. Avec mois_debut=1, l'exercice est l'année civile."""
    return d.year if d.month >= mois_debut else d.year - 1


def nouvel_id() -> str:
    return uuid.uuid4().hex[:12]


def nouvelle_operation(d: date, manifestation: str, sens: str, categorie: str, libelle: str,
                       montant: float, mode: str, tiers: str, piece: str, utilisateur: str,
                       mois_debut: int = 1) -> dict:
    return {
        "id": nouvel_id(), "date": d.isoformat(), "exercice": exercice_de(d, mois_debut),
        "manifestation": manifestation, "sens": sens, "categorie": categorie,
        "libelle": libelle.strip(), "montant": round(abs(float(montant)), 2), "mode_paiement": mode,
        "tiers": tiers.strip(), "piece": piece.strip(), "saisi_par": utilisateur,
        "saisi_le": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


# ------------------------------------------------------------------ synthèses
def synthese_manifestations(ops: pd.DataFrame, exercice: int) -> pd.DataFrame:
    x = ops[ops["exercice"] == exercice]
    if x.empty:
        return pd.DataFrame(columns=["manifestation", "Recettes", "Dépenses", "Résultat"])
    t = x.pivot_table(index="manifestation", columns="sens", values="montant", aggfunc="sum", fill_value=0)
    t = t.reindex(columns=SENS, fill_value=0).rename(columns={"Recette": "Recettes", "Dépense": "Dépenses"})
    t["Résultat"] = t["Recettes"] - t["Dépenses"]
    return t.reset_index().sort_values("Résultat", ascending=False)


def synthese_categories(ops: pd.DataFrame, exercice: int) -> pd.DataFrame:
    x = ops[ops["exercice"] == exercice]
    return (x.groupby(["sens", "categorie"], as_index=False)["montant"].sum()
             .sort_values(["sens", "categorie"], ascending=[False, True]))


# ------------------------------------------------------------------ modèle d'import
def modele_import_excel() -> bytes:
    exemple = pd.DataFrame([
        ["15/05/2026", "Brocante", "Recette", "708", "Emplacements brocante", 640.00, "Espèces", "", "B-001"],
        ["15/05/2026", "Brocante", "Dépense", "607", "Boissons buvette", 312.40, "Carte bancaire", "Grossiste", "F-118"],
        ["03/02/2026", MANIFESTATION_GENERALE, "Recette", "Subventions", "Subvention mairie", 1500, "Virement", "Mairie", ""],
    ], columns=COLONNES_IMPORT)
    aide = pd.DataFrame({
        "Colonne": COLONNES_IMPORT,
        "Obligatoire": ["Oui", "Oui", "Oui*", "Non", "Oui", "Oui", "Non", "Non", "Non"],
        "Ce qu'il faut mettre": [
            "Date de l'opération (JJ/MM/AAAA)",
            "Nom de la manifestation ; si elle n'existe pas, elle peut être créée à l'import",
            "Recette ou Dépense (*facultatif si le montant est négatif pour une dépense)",
            "Code (ex. 708) ou nom de catégorie ; vide = « À classer »",
            "Description courte de l'opération",
            "Montant en euros (virgule ou point acceptés)",
            "Espèces, Chèque, Virement, Carte bancaire, Autre",
            "Fournisseur, donateur, prestataire…",
            "Numéro de facture, de ticket ou de reçu",
        ],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        exemple.to_excel(w, sheet_name="Opérations", index=False)
        aide.to_excel(w, sheet_name="Mode d'emploi", index=False)
        _mettre_en_forme(w)
    return buf.getvalue()


# ------------------------------------------------------------------ import contrôlé
_SYNONYMES = {
    "date": "Date", "manifestation": "Manifestation", "évènement": "Manifestation", "evenement": "Manifestation",
    "événement": "Manifestation", "activité": "Manifestation", "sens": "Sens", "type": "Sens",
    "catégorie": "Catégorie", "categorie": "Catégorie", "compte": "Catégorie", "libellé": "Libellé",
    "libelle": "Libellé", "désignation": "Libellé", "objet": "Libellé", "montant": "Montant",
    "mode de paiement": "Mode de paiement", "paiement": "Mode de paiement", "règlement": "Mode de paiement",
    "tiers": "Tiers", "fournisseur": "Tiers", "n° pièce": "N° pièce", "pièce": "N° pièce", "piece": "N° pièce",
}


def empreinte(d, manif, montant, libelle) -> str:
    cle = f"{pd.Timestamp(d).date()}|{str(manif).strip().lower()}|{round(float(montant), 2)}|{str(libelle).strip().lower()}"
    return hashlib.sha1(cle.encode()).hexdigest()


def analyser_import(fichier, manifestations_connues: list[str], ops_existantes: pd.DataFrame,
                    mois_debut: int = 1) -> tuple[pd.DataFrame, list[str]]:
    """Lit le fichier, normalise chaque ligne et renvoie (aperçu, erreurs globales).

    L'aperçu contient une colonne « statut » : OK, Avertissement, Doublon ou Erreur, et « message ».
    """
    try:
        brut = pd.read_excel(fichier, sheet_name=0)
    except Exception as e:  # fichier illisible
        return pd.DataFrame(), [f"Fichier illisible : {e}"]
    brut = brut.rename(columns=lambda c: _SYNONYMES.get(str(c).strip().lower(), str(c).strip()))
    manquantes = [c for c in ("Date", "Libellé", "Montant") if c not in brut.columns]
    if manquantes:
        return pd.DataFrame(), [f"Colonnes obligatoires absentes : {', '.join(manquantes)}. "
                                f"Utilisez le modèle d'import."]
    brut = brut.dropna(how="all")
    idx_cat = index_categories()
    connues = {m.strip().lower(): m for m in manifestations_connues}
    modes = {m.lower(): m for m in MODES_PAIEMENT}
    existantes = set()
    if not ops_existantes.empty:
        for r in ops_existantes.itertuples():
            if pd.notna(r.date):
                existantes.add(empreinte(r.date, r.manifestation, r.montant, r.libelle))

    lignes, vues = [], set()
    for n, r in enumerate(brut.to_dict("records"), start=2):
        msgs, statut = [], "OK"
        d = pd.to_datetime(r.get("Date"), dayfirst=True, errors="coerce")
        m = pd.to_numeric(str(r.get("Montant", "")).replace("€", "").replace(" ", "").replace(",", "."),
                          errors="coerce")
        libelle = "" if pd.isna(r.get("Libellé")) else str(r.get("Libellé")).strip()
        if pd.isna(d):
            msgs.append("date illisible"); statut = "Erreur"
        if pd.isna(m) or m == 0:
            msgs.append("montant absent ou nul"); statut = "Erreur"
        if not libelle:
            msgs.append("libellé vide"); statut = "Erreur"

        sens_brut = str(r.get("Sens", "") if pd.notna(r.get("Sens", "")) else "").strip().lower()
        if sens_brut.startswith("r") or sens_brut in ("+", "entrée", "entree", "crédit", "credit"):
            sens = "Recette"
        elif sens_brut.startswith("d") or sens_brut in ("-", "sortie", "débit", "debit"):
            sens = "Dépense"
        elif pd.notna(m) and m < 0:
            sens = "Dépense"
        else:
            sens = None

        cat_brut = r.get("Catégorie", "")
        cat_brut = "" if pd.isna(cat_brut) else str(cat_brut).strip()
        trouve = idx_cat.get(cat_brut.lower()) if cat_brut else None
        if trouve:
            if sens is None:
                sens = trouve[0]
            categorie = trouve[1] if trouve[0] == sens else None
            if categorie is None:
                msgs.append(f"catégorie « {cat_brut} » incompatible avec {sens}, mise à classer")
                statut = _pire(statut, "Avertissement")
        else:
            categorie = None
            if cat_brut:
                msgs.append(f"catégorie « {cat_brut} » inconnue, mise à classer")
                statut = _pire(statut, "Avertissement")
        if sens is None:
            msgs.append("sens inconnu (Recette/Dépense)"); statut = "Erreur"; sens = ""
        if categorie is None and sens:
            categorie = categorie_a_classer(sens)

        manif_brut = r.get("Manifestation", "")
        manif_brut = "" if pd.isna(manif_brut) else str(manif_brut).strip()
        if not manif_brut:
            manif = MANIFESTATION_GENERALE
        elif manif_brut.lower() in connues:
            manif = connues[manif_brut.lower()]
        else:
            manif = manif_brut
            msgs.append(f"nouvelle manifestation « {manif} »")
            statut = _pire(statut, "Avertissement")

        mode_brut = "" if pd.isna(r.get("Mode de paiement", "")) else str(r.get("Mode de paiement", "")).strip()
        mode = modes.get(mode_brut.lower(), "Autre" if mode_brut else "")

        if statut != "Erreur":
            emp = empreinte(d, manif, abs(m), libelle)
            if emp in existantes:
                statut, msgs = "Doublon", msgs + ["déjà présente dans la base"]
            elif emp in vues:
                statut, msgs = "Doublon", msgs + ["en double dans le fichier"]
            vues.add(emp)

        lignes.append({
            "ligne": n, "statut": statut, "message": " ; ".join(msgs),
            "date": None if pd.isna(d) else d.date(), "exercice": None if pd.isna(d) else exercice_de(d, mois_debut),
            "manifestation": manif, "sens": sens, "categorie": categorie or "", "libelle": libelle,
            "montant": None if pd.isna(m) else round(abs(float(m)), 2), "mode_paiement": mode,
            "tiers": "" if pd.isna(r.get("Tiers", "")) else str(r.get("Tiers", "")).strip(),
            "piece": "" if pd.isna(r.get("N° pièce", "")) else str(r.get("N° pièce", "")).strip(),
        })
    return pd.DataFrame(lignes), []


def _pire(a: str, b: str) -> str:
    ordre = ["OK", "Avertissement", "Doublon", "Erreur"]
    return a if ordre.index(a) >= ordre.index(b) else b


def lignes_a_importer(apercu: pd.DataFrame, inclure_doublons: bool, utilisateur: str) -> list[dict]:
    statuts = {"OK", "Avertissement"} | ({"Doublon"} if inclure_doublons else set())
    ok = apercu[apercu["statut"].isin(statuts)]
    horodatage = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return [{
        "id": nouvel_id(), "date": r["date"].isoformat(), "exercice": int(r["exercice"]),
        "manifestation": r["manifestation"], "sens": r["sens"], "categorie": r["categorie"],
        "libelle": r["libelle"], "montant": float(r["montant"]), "mode_paiement": r["mode_paiement"],
        "tiers": r["tiers"], "piece": r["piece"], "saisi_par": f"import ({utilisateur})", "saisi_le": horodatage,
    } for r in ok.to_dict("records")]


# ------------------------------------------------------------------ sauvegarde Excel
def export_excel(tables: dict[str, pd.DataFrame]) -> bytes:
    ops = typer_operations(tables["operations"])
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        vue = ops.sort_values("date").drop(columns=["signe"])
        vue["date"] = vue["date"].dt.date
        vue.rename(columns={"id": "Identifiant", "date": "Date", "exercice": "Exercice",
                            "manifestation": "Manifestation", "sens": "Sens", "categorie": "Catégorie",
                            "libelle": "Libellé", "montant": "Montant", "mode_paiement": "Mode de paiement",
                            "tiers": "Tiers", "piece": "N° pièce", "saisi_par": "Saisi par",
                            "saisi_le": "Saisi le"}).to_excel(w, sheet_name="Opérations", index=False)
        lignes = []
        for ex in sorted(ops["exercice"].unique()):
            s = synthese_manifestations(ops, ex)
            s.insert(0, "Exercice", ex)
            lignes.append(s)
        if lignes:
            pd.concat(lignes).to_excel(w, sheet_name="Synthèse", index=False)
        for t, df in tables.items():
            if t != "operations":
                df.to_excel(w, sheet_name=t.replace("_", " ").capitalize()[:31], index=False)
        _mettre_en_forme(w)
    return buf.getvalue()


def _mettre_en_forme(writer) -> None:
    from openpyxl.styles import Font, PatternFill
    for ws in writer.book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        ws.freeze_panes = "A2"
        for col in ws.columns:
            largeur = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, largeur + 2), 60)
            if str(col[0].value).lower() in ("montant", "recettes", "dépenses", "résultat", "depenses"):
                for c in col[1:]:
                    c.number_format = '#,##0.00 "€"'


# ------------------------------------------------------------------ prévisionnel
def base_previsionnelle(ops: pd.DataFrame, manifestations: list[str], annee_cible: int,
                        methode: str, hausse_depenses: float, evolution_recettes: float) -> pd.DataFrame:
    """Propose recettes/dépenses N+1 par manifestation à partir de l'historique.

    methode : « derniere » (dernier exercice réalisé) ou « moyenne3 » (moyenne des 3 derniers exercices
    où la manifestation a eu lieu).
    """
    passe = ops[ops["exercice"] < annee_cible]
    lignes = []
    for manif in manifestations:
        x = passe[passe["manifestation"] == manif]
        annees = sorted(x["exercice"].unique(), reverse=True)
        if methode == "derniere":
            annees = annees[:1]
        else:
            annees = annees[:3]
        x = x[x["exercice"].isin(annees)]
        n = max(len(annees), 1)
        rec = x.loc[x["sens"] == "Recette", "montant"].sum() / n
        dep = x.loc[x["sens"] == "Dépense", "montant"].sum() / n
        lignes.append({
            "manifestation": manif,
            "historique": ", ".join(str(a) for a in sorted(annees)) or "aucun",
            "recettes": round(rec * (1 + evolution_recettes / 100), 2),
            "depenses": round(dep * (1 + hausse_depenses / 100), 2),
        })
    df = pd.DataFrame(lignes)
    if not df.empty:
        df["resultat"] = df["recettes"] - df["depenses"]
    return df
